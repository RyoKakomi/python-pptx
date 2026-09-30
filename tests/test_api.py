"""Unit-test suite for `pptx.api` module."""

from __future__ import annotations

import os

import pytest

import pptx
from pptx.api import Presentation
from pptx.opc.constants import CONTENT_TYPE as CT
from pptx.parts.presentation import PresentationPart

from .unitutil.mock import class_mock, instance_mock


class DescribePresentation(object):
    def it_opens_default_template_on_no_path_provided(self, call_fixture):
        Package_, path, prs_ = call_fixture
        prs = Presentation()
        Package_.open.assert_called_once_with(path)
        assert prs is prs_

    # fixtures -------------------------------------------------------

    @pytest.fixture
    def call_fixture(self, Package_, prs_, prs_part_):
        path = os.path.abspath(
            os.path.join(os.path.split(pptx.__file__)[0], "templates", "default.pptx")
        )
        Package_.open.return_value.main_document_part = prs_part_
        prs_part_.content_type = CT.PML_PRESENTATION_MAIN
        prs_part_.presentation = prs_
        return Package_, path, prs_

    # fixture components ---------------------------------------------

    @pytest.fixture
    def Package_(self, request):
        return class_mock(request, "pptx.api.Package")

    @pytest.fixture
    def prs_(self, request):
        return instance_mock(request, Presentation)

    @pytest.fixture
    def prs_part_(self, request):
        return instance_mock(request, PresentationPart)


class DescribeDefaultTemplate(object):
    """Integration tests for the core properties shipped in the default template."""

    def it_has_no_personal_name_in_core_properties(self):
        core_props = Presentation().core_properties

        assert core_props.author == ""
        assert core_props.last_modified_by == "python-pptx"
        assert "Steve Canny" not in core_props.blob.decode("utf-8")
